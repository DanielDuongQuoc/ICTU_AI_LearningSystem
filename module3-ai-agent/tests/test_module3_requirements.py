"""Module 3 requirements check: Word (.docx) support, exam-matrix generation, scoring/fraud separation (offline)."""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Optional

import pytest

from app.core.config import get_settings
from app.core.exceptions import AppError, ErrorCode
from app.schemas.evaluation import EvaluationResponse
from app.services import docx_service
from tests.fakes import RecordingLLM, Stack

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def _settings(monkeypatch):
    monkeypatch.setenv("RATE_LIMIT_PER_MINUTE", "0")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def make_docx(blocks: list[Any]) -> bytes:
    """blocks: str paragraph | ("table", [[cells]]) | "<PAGEBREAK>" | ("heading", text)."""
    import docx

    d = docx.Document()
    for b in blocks:
        if b == "<PAGEBREAK>":
            d.add_page_break()
        elif isinstance(b, tuple) and b[0] == "table":
            t = d.add_table(rows=len(b[1]), cols=len(b[1][0]))
            for r, row in enumerate(b[1]):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
        elif isinstance(b, tuple) and b[0] == "heading":
            d.add_heading(b[1], level=1)
        else:
            d.add_paragraph(b)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def upload(client, data: bytes, name: str = "lecture.docx", mime: str = DOCX_MIME):
    return await client.post("/api/documents/upload", files={"file": (name, data, mime)})


# ===================================================================== DOCX
@pytest.mark.asyncio
async def test_docx_upload_is_indexed_and_retrievable_with_logical_pages(client, monkeypatch):
    Stack().install(monkeypatch)
    data = make_docx([
        ("heading", "Chapter one"),
        " ".join(f"Zebracorn systems explain machine learning topic {i}. Zebracorn models learn patterns from data." for i in range(14)),
        "<PAGEBREAK>",
        "Quokkaplex indexes speed database queries. Quokkaplex structures reduce lookup time.",
        ("table", [["Narwhalon", "cache"], ["Hedgeflux", "queue"]]),
    ])
    r = await upload(client, data)
    assert r.status_code in (200, 201), r.text
    body = r.json()
    assert body["filename"] == "lecture.docx" and body["status"] == "indexed"
    assert body["pages"] == 2 and body["chunks"] >= 2
    doc = body["document_id"]

    hit = await client.post("/api/rag/retrieve", json={"document_id": doc, "question": "quokkaplex indexes database queries"})
    assert hit.status_code == 200, hit.text
    chunks = hit.json()["chunks"]
    assert chunks and all(c["document_id"] == doc for c in chunks)
    # A chunk that straddles a page break reports its first page (same rule as PDF), so check the exact page at extraction level.
    assert any("Quokkaplex" in c["text"] and c["page"] in (1, 2) for c in chunks), chunks
    assert all("filename" not in c or c["filename"] for c in chunks)
    tab = await client.post("/api/rag/retrieve", json={"document_id": doc, "question": "narwhalon cache hedgeflux queue"})
    assert any("Narwhalon | cache" in c["text"] for c in tab.json()["chunks"]), "table rows must be extracted"

    assert (await client.delete(f"/api/documents/{doc}")).status_code == 200
    assert (await client.get(f"/api/documents/{doc}")).status_code == 404


@pytest.mark.asyncio
async def test_docx_and_pdf_documents_stay_isolated(client, monkeypatch):
    from tests.fakes import sample_pdf_bytes

    Stack().install(monkeypatch)
    d = await upload(client, make_docx(["Zebracorn systems explain machine learning from data."]))
    p = await client.post("/api/documents/upload", files={"file": ("a.pdf", sample_pdf_bytes(["Quokkaplex indexes speed database queries."]), "application/pdf")})
    assert d.status_code in (200, 201) and p.status_code in (200, 201)
    hit = await client.post("/api/rag/retrieve", json={"document_id": d.json()["document_id"], "question": "quokkaplex indexes"})
    assert all("Quokkaplex" not in c["text"] for c in hit.json()["chunks"])


def test_page_breaks_define_logical_pages(tmp_path):
    path = tmp_path / "pb.docx"
    path.write_bytes(make_docx(["First page text.", "<PAGEBREAK>", "Second page text.", ("table", [["a", "b"]]), "<PAGEBREAK>", "Third."]))
    pages = docx_service.extract_pages(path)
    assert [(p.page, p.text) for p in pages] == [(1, "First page text."), (2, "Second page text.\n\na | b"), (3, "Third.")]


def test_long_page_is_split_into_logical_pages(tmp_path):
    path = tmp_path / "long.docx"
    path.write_bytes(make_docx([f"Paragraph {i}: " + "word " * 40 for i in range(60)]))
    pages = docx_service.extract_pages(path)
    assert len(pages) > 1
    assert [p.page for p in pages] == list(range(1, len(pages) + 1))
    assert all(len(p.text) <= docx_service.LOGICAL_PAGE_CHARS + 400 for p in pages)


def test_blank_docx_has_no_extractable_text(tmp_path):
    path = tmp_path / "blank.docx"
    path.write_bytes(make_docx([]))
    with pytest.raises(AppError) as e:
        docx_service.extract_pages(path)
    assert e.value.code == ErrorCode.NO_EXTRACTABLE_TEXT


@pytest.mark.asyncio
async def test_docx_content_is_checked_not_just_the_name(client, monkeypatch):
    Stack().install(monkeypatch)
    for data in (b"plain text pretending to be a document", b"PK\x03\x04 not really a zip"):
        r = await upload(client, data)
        assert r.status_code == 400, r.text
        assert r.json()["error"]["code"] == "INVALID_FILE"


@pytest.mark.asyncio
async def test_other_formats_and_wrong_mime_are_rejected(client, monkeypatch):
    Stack().install(monkeypatch)
    ok = make_docx(["Zebracorn systems explain machine learning from data."])
    assert (await upload(client, ok, "x.docx", "image/png")).status_code == 400   # wrong MIME
    assert (await upload(client, ok, "x.doc", "application/msword")).status_code == 400   # legacy Word
    assert (await upload(client, ok, "x.docm", "application/vnd.ms-word.document.macroEnabled.12")).status_code == 400


def _zip(names_to_data: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in names_to_data.items():
            z.writestr(n, d)
    return buf.getvalue()


def test_macro_zip_bomb_and_non_word_zip_are_rejected(tmp_path):
    base = {"[Content_Types].xml": b"<x/>", "word/document.xml": b"<x/>"}
    cases = {
        "macro": _zip({**base, "word/vbaProject.bin": b"x"}),
        "notword": _zip({"hello.txt": b"hi"}),
        "bomb": _zip({**base, "word/big.xml": b"0" * 3_000_000}),
    }
    for name, data in cases.items():
        p = tmp_path / f"{name}.docx"
        p.write_bytes(data)
        with pytest.raises(AppError) as e:
            docx_service.validate_docx_path(p, max_bytes=200_000, declared_filename=p.name)
        assert e.value.code in (ErrorCode.INVALID_FILE, ErrorCode.DOCUMENT_LIMIT_EXCEEDED), name


@pytest.mark.asyncio
async def test_oversized_docx_is_rejected(client, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_SIZE_MB", "1")
    get_settings.cache_clear()
    Stack().install(monkeypatch)
    big = make_docx(["x"]) + b"0" * (1024 * 1024 + 10)
    assert (await upload(client, big)).status_code == 413


# ===================================================================== exam matrix
DOC_TEXT = " ".join([
    "Quokkaplex indexes speed database queries.", "Narwhalon caches store frequent results.",
    "Hedgeflux queues order background jobs.", "Zebracorn models learn patterns from data.",
    "Pangolith schedulers balance worker load.", "Marmotron logs record every request.",
    "Lemuroid routers forward network packets.", "Okapiton filters remove invalid input.",
])
TOPICS = ["quokkaplex indexes speed database queries", "narwhalon caches store frequent results",
          "hedgeflux queues order background jobs", "zebracorn models learn patterns from data",
          "pangolith schedulers balance worker load", "marmotron logs record every request",
          "lemuroid routers forward network packets", "okapiton filters remove invalid input"]


class MatrixLLM(RecordingLLM):
    """Honours the requested type/difficulty and yields a fresh grounded question each time."""

    def __init__(self, per_call: int = 6, fail_from_call: Optional[int] = None, same_text: bool = False) -> None:
        super().__init__()
        self.per_call, self.fail_from_call, self.same_text = per_call, fail_from_call, same_text
        self.gen_calls = 0
        self.counter = 0

    async def generate_json(self, system_prompt: str, user_prompt: str, *, required_keys=None, **kw: Any) -> Any:
        if "questions" not in (required_keys or []):
            return await super().generate_json(system_prompt, user_prompt, required_keys=required_keys, **kw)
        self.calls.append(("json", system_prompt, user_prompt))
        self.gen_calls += 1
        if self.fail_from_call is not None and self.gen_calls >= self.fail_from_call:
            raise AppError(ErrorCode.OLLAMA_TIMEOUT, "model timed out")
        diff = re.search(r"Difficulty: (\w+)", user_prompt).group(1)
        qtype = re.search(r"Question type: (\w+)", user_prompt).group(1)
        out = []
        for _ in range(self.per_call):
            topic = TOPICS[0 if self.same_text else self.counter % len(TOPICS)]
            self.counter += 1
            q = {"question": f"What does the document explain about {topic}?",
                 "expected_answer": f"It states that {topic}.", "difficulty": diff, "type": qtype}
            if qtype == "multiple_choice":
                opts = [f"{topic}", f"Option B{self.counter}", f"Option C{self.counter}", f"Option D{self.counter}"]
                q.update(options=opts, correct_answer=opts[0])
            elif qtype == "true_false":
                q.update(correct_answer="true", options=["True", "False"])
            out.append(q)
        return {"questions": out}


async def matrix(client, doc, items, subject=None):
    body = {"document_id": doc, "items": items}
    if subject is not None:
        body["subject"] = subject
    return await client.post("/api/questions/generate-matrix", json=body)


async def indexed_doc(client, monkeypatch, llm):
    stack = Stack(llm).install(monkeypatch)
    r = await upload(client, make_docx([DOC_TEXT]))
    assert r.status_code in (200, 201), r.text
    return stack, r.json()["document_id"]


@pytest.mark.asyncio
async def test_matrix_generates_exact_counts_types_and_difficulties(client, monkeypatch):
    llm = MatrixLLM()
    _, doc = await indexed_doc(client, monkeypatch, llm)
    items = [{"question_type": "multiple_choice", "difficulty": "medium", "count": 2},
             {"question_type": "true_false", "difficulty": "easy", "count": 1},
             {"question_type": "short_answer", "difficulty": "hard", "count": 2}]
    r = await matrix(client, doc, items, subject="  Databases   101 ")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["subject"] == "Databases 101" and body["requested_total"] == 5 == body["generated_total"]
    assert body["complete"] is True
    assert [i["status"] for i in body["items"]] == ["ok"] * 3
    qs = body["questions"]
    assert len(qs) == 5 and len({q["question_id"] for q in qs}) == 5
    for idx, it in enumerate(items):
        mine = [q for q in qs if q["matrix_item"] == idx]
        assert len(mine) == it["count"]
        assert all(q["type"] == it["question_type"] and q["difficulty"] == it["difficulty"] for q in mine)
    assert all(q["source_chunks"] for q in qs), "every question is traceable to document chunks"
    assert all(len(q["options"]) == 4 and q["correct_answer"] in q["options"] for q in qs if q["type"] == "multiple_choice")
    assert len({q["question"] for q in qs}) == 5, "no duplicate across rows"
    assert "Databases" not in llm.prompts, "subject is not sent to the model"
    assert llm.gen_calls == 3, "one gated LLM call per matrix row"


@pytest.mark.asyncio
async def test_matrix_reports_shortfall_instead_of_padding(client, monkeypatch):
    llm = MatrixLLM(per_call=1)
    _, doc = await indexed_doc(client, monkeypatch, llm)
    r = await matrix(client, doc, [{"question_type": "short_answer", "difficulty": "easy", "count": 3},
                                    {"question_type": "short_answer", "difficulty": "hard", "count": 1}])
    body = r.json()
    assert r.status_code == 200
    assert body["complete"] is False and body["requested_total"] == 4 and body["generated_total"] == 2
    assert [(i["requested"], i["generated"], i["status"]) for i in body["items"]] == [(3, 1, "partial"), (1, 1, "ok")]


@pytest.mark.asyncio
async def test_matrix_later_row_failure_is_reported_and_earlier_rows_kept(client, monkeypatch):
    llm = MatrixLLM(fail_from_call=2)
    _, doc = await indexed_doc(client, monkeypatch, llm)
    r = await matrix(client, doc, [{"question_type": "short_answer", "difficulty": "easy", "count": 2},
                                    {"question_type": "essay", "difficulty": "hard", "count": 2}])
    body = r.json()
    assert r.status_code == 200 and body["complete"] is False
    assert body["items"][0]["status"] == "ok"
    assert body["items"][1]["status"] == "failed" and body["items"][1]["error_code"] == "OLLAMA_TIMEOUT"
    assert body["generated_total"] == 2 and all(q["matrix_item"] == 0 for q in body["questions"])


@pytest.mark.asyncio
async def test_matrix_first_row_failure_fails_the_request(client, monkeypatch):
    llm = MatrixLLM(fail_from_call=1)
    _, doc = await indexed_doc(client, monkeypatch, llm)
    r = await matrix(client, doc, [{"question_type": "short_answer", "difficulty": "easy", "count": 2}])
    assert r.status_code == 504 and r.json()["error"]["code"] == "OLLAMA_TIMEOUT"


@pytest.mark.asyncio
async def test_matrix_removes_duplicates_across_rows(client, monkeypatch):
    llm = MatrixLLM(per_call=1, same_text=True)
    _, doc = await indexed_doc(client, monkeypatch, llm)
    r = await matrix(client, doc, [{"question_type": "short_answer", "difficulty": "easy", "count": 1},
                                    {"question_type": "short_answer", "difficulty": "medium", "count": 1}])
    body = r.json()
    assert body["generated_total"] == 1 and body["complete"] is False
    assert body["items"][1]["status"] == "failed"


@pytest.mark.asyncio
async def test_matrix_unknown_document_is_404(client, monkeypatch):
    Stack(MatrixLLM()).install(monkeypatch)
    r = await matrix(client, "11111111-1111-1111-1111-111111111111",
                     [{"question_type": "short_answer", "difficulty": "easy", "count": 1}])
    assert r.status_code == 404


@pytest.mark.asyncio
@pytest.mark.parametrize("items", [
    [],
    [{"question_type": "short_answer", "difficulty": "easy", "count": 0}],
    [{"question_type": "short_answer", "difficulty": "easy", "count": 21}],
    [{"question_type": "poem", "difficulty": "easy", "count": 1}],
    [{"question_type": "short_answer", "difficulty": "impossible", "count": 1}],
    [{"question_type": "short_answer", "difficulty": "easy", "count": 20}] * 3,                    # 60 > 50
    [{"question_type": "short_answer", "difficulty": "easy", "count": 1}] * 13,                    # > 12 rows
])
async def test_matrix_invalid_requests_are_422(client, monkeypatch, items):
    Stack(MatrixLLM()).install(monkeypatch)
    r = await matrix(client, "11111111-1111-1111-1111-111111111111", items)
    assert r.status_code == 422, r.text


@pytest.mark.asyncio
async def test_matrix_total_limit_is_configurable(client, monkeypatch):
    monkeypatch.setenv("MAX_MATRIX_QUESTIONS", "3")
    get_settings.cache_clear()
    Stack(MatrixLLM()).install(monkeypatch)
    r = await matrix(client, "11111111-1111-1111-1111-111111111111",
                     [{"question_type": "short_answer", "difficulty": "easy", "count": 4}])
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_existing_single_type_generation_is_unchanged(client, monkeypatch):
    llm = MatrixLLM()
    _, doc = await indexed_doc(client, monkeypatch, llm)
    r = await client.post("/api/questions/generate", json={"document_id": doc, "number_of_questions": 2,
                                                             "difficulty": "easy", "question_type": "short_answer"})
    assert r.status_code == 200 and len(r.json()["questions"]) == 2
    assert "matrix_item" not in r.json()["questions"][0]


# ===================================================================== scoring is not fraud detection
def test_evaluation_contract_has_no_fraud_or_cheating_fields():
    assert set(EvaluationResponse.model_fields) == {
        "score", "max_score", "percentage", "correct", "feedback", "strengths", "missing_points", "errors", "sources"}


def test_no_fraud_api_exists():
    from app.main import app

    paths = " ".join(r.path for r in app.routes if hasattr(r, "path")).lower()
    assert not re.search(r"fraud|cheat|plagiar|anti", paths)
